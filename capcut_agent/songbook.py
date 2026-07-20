"""송북(Songbook): 곡별 정답 가사 · 무드 · 배경 가이드.

Mindtrack 스타일의 엑셀(`가사 원문`, `분위기 가이드` 시트)에서 곡 정보를
읽어옵니다. 정답 가사는 Whisper 자동 자막을 교정하는 기준으로 쓰입니다.

가사 클린업(`clean_lyric_lines`)과 무드 매핑은 순수 함수라 엑셀 없이도
테스트됩니다. 엑셀 파싱(`load_songbook`)만 openpyxl 을 지연 임포트합니다.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional

# 6개 무드 태그 (분위기 가이드 시트 기준)
MOODS = ("각성", "버팀", "확신", "도약", "도착", "위로")

# 무드 → 전환/자막 스타일 프리셋 (transitions.PRESETS)
MOOD_STYLE: Dict[str, str] = {
    "각성": "goosebump",   # 시작·결심·자극, 딥 블루, 비트 드롭
    "버팀": "dreamy",      # 인내·꾸준함, 차분한 틸
    "확신": "goosebump",   # 자기신뢰·시각화, 앰버 빌드업
    "도약": "energetic",   # 추진·질주·올인, 속도감
    "도착": "goosebump",   # 성취·완주, 앤섬
    "위로": "dreamy",      # 여운·다독임, 뮤트 그레이
}

# 무드 → 배경 영상/색감 가이드 (클립 고를 때 참고용)
MOOD_BACKGROUND: Dict[str, Dict[str, str]] = {
    "각성": {"emotion": "시작·결심·자극", "video": "새벽 도심 야경, 출발하는 차", "color": "딥 블루"},
    "버팀": {"emotion": "인내·꾸준함", "video": "빗길, 터널, 가로등 도로", "color": "차분한 틸"},
    "확신": {"emotion": "자기신뢰·시각화", "video": "해 뜨기 직전 도로, 한강뷰", "color": "앰버 포인트"},
    "도약": {"emotion": "추진·질주·올인", "video": "트인 고속도로, 속도감 POV", "color": "밝은 블루-화이트"},
    "도착": {"emotion": "성취·완주", "video": "정상 뷰, 밝아오는 하늘", "color": "따뜻한 골드"},
    "위로": {"emotion": "여운·다독임", "video": "정차한 차 안, 비 그친 거리", "color": "뮤트 그레이"},
}

DEFAULT_STYLE = "goosebump"


@dataclass
class Song:
    """한 곡의 메타데이터 + 정답 가사.

    Attributes:
        title: 곡명.
        mood: 무드 태그(MOODS 중 하나).
        has_file: 원본 오디오 파일 보유 여부.
        duration: 길이(초). 알 수 없으면 None.
        link: 공개 링크(Suno 등).
        raw_lyrics: 섹션 태그가 포함된 원문 가사.
    """

    title: str
    mood: str = ""
    has_file: bool = False
    duration: Optional[float] = None
    link: str = ""
    raw_lyrics: str = ""

    def lyric_lines(self, keep_adlibs: bool = True) -> List[str]:
        """부를 수 있는 가사 줄 목록(섹션 태그 제거)."""
        return clean_lyric_lines(self.raw_lyrics, keep_adlibs=keep_adlibs)

    def style(self) -> str:
        """무드에 대응하는 스타일 프리셋 이름."""
        return mood_to_style(self.mood)

    def background_guide(self) -> Dict[str, str]:
        return MOOD_BACKGROUND.get(self.mood, {})


def mood_to_style(mood: str) -> str:
    """무드 태그 → 스타일 프리셋 이름. 모르면 기본값."""
    return MOOD_STYLE.get((mood or "").strip(), DEFAULT_STYLE)


_SECTION_RE = re.compile(r"^\[.*\]$")


def clean_lyric_lines(raw: str, *, keep_adlibs: bool = True) -> List[str]:
    """원문 가사에서 자막용 '부를 수 있는 줄'만 추립니다 (순수 함수).

    - `[Intro]`, `[Chorus, ...]` 같은 섹션 태그 줄은 제거.
    - 빈 줄 제거.
    - 줄 끝에 남은 인라인 섹션 태그도 정리.
    - `keep_adlibs=False` 면 `(...)` 백보컬/애드립 전용 줄도 제거.

    Args:
        raw: 원문 가사(줄바꿈 포함).
        keep_adlibs: 괄호 애드립 줄 유지 여부.

    Returns:
        가사 줄 리스트.
    """
    lines: List[str] = []
    for line in (raw or "").splitlines():
        s = line.strip()
        if not s:
            continue
        if _SECTION_RE.match(s):  # [Verse 1] 등 섹션 태그 줄
            continue
        # 줄 안에 인라인으로 붙은 [..] 태그 제거 (드묾)
        s = re.sub(r"\[[^\]]*\]", "", s).strip()
        if not s:
            continue
        if not keep_adlibs and s.startswith("(") and s.endswith(")"):
            continue
        lines.append(s)
    return lines


def _parse_duration(text: str) -> Optional[float]:
    """'1:54' / '3:39' → 초. 파싱 실패 시 None."""
    if not text:
        return None
    parts = str(text).strip().split(":")
    try:
        nums = [int(p) for p in parts]
    except ValueError:
        return None
    if len(nums) == 2:
        return float(nums[0] * 60 + nums[1])
    if len(nums) == 3:
        return float(nums[0] * 3600 + nums[1] * 60 + nums[2])
    if len(nums) == 1:
        return float(nums[0])
    return None


def load_songbook(xlsx_path: str, *, sheet: str = "가사 원문") -> Dict[str, Song]:
    """엑셀에서 곡별 정답 가사를 읽어 {곡명: Song} 으로 반환 (openpyxl 지연 임포트).

    `가사 원문` 시트 형식(헤더: 곡명|분위기|파일|길이|공개 링크|가사 원문)을 가정하되,
    무드 태그가 있는 데이터 행만 취합니다(섹션 구분 행은 건너뜀).
    """
    import openpyxl  # noqa: WPS433

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise KeyError(f"시트 '{sheet}' 가 없습니다. 사용 가능: {wb.sheetnames}")
    ws = wb[sheet]

    songs: Dict[str, Song] = {}
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 6:
            continue
        title = (str(row[0]).strip() if row[0] is not None else "")
        mood = (str(row[1]).strip() if row[1] is not None else "")
        lyrics = row[5] if len(row) > 5 else None
        if not title or mood not in MOODS or not lyrics:
            continue  # 헤더/구분 행 제외
        songs[title] = Song(
            title=title,
            mood=mood,
            has_file=(str(row[2]).strip() == "있음") if row[2] is not None else False,
            duration=_parse_duration(str(row[3]) if row[3] is not None else ""),
            link=str(row[4]).strip() if row[4] is not None else "",
            raw_lyrics=str(lyrics),
        )
    return songs


def find_song(songs: Dict[str, Song], name: str) -> Optional[Song]:
    """곡명으로 곡을 찾습니다(대소문자·공백 무시, 부분일치 허용)."""
    if not name:
        return None
    if name in songs:
        return songs[name]
    norm = _norm_title(name)
    # 정확 일치(정규화)
    for song in songs.values():
        if _norm_title(song.title) == norm:
            return song
    # 부분 일치
    matches = [s for s in songs.values() if norm in _norm_title(s.title)]
    return matches[0] if len(matches) == 1 else (matches[0] if matches else None)


def _norm_title(t: str) -> str:
    return re.sub(r"\s+", " ", (t or "").strip().lower())
