"""벤치마킹 보드 — 분석 결과를 엑셀 한 파일로.

    from threadscout.benchmark import write_xlsx
    write_xlsx(result, "output/benchmark.xlsx")

시트 구성
    1_벤치마킹      상위 글 한 줄씩 — 점수·노출·댓글율·공유율·훅 유형·따라할 포인트
    2_훅패턴        훅/길이/미디어/시간대별 성과 (무엇을 따라할지)
    3_키워드        상위권에서 유독 많이 나온 단어·해시태그
    4_계정          벤치마킹할 계정
    5_쿠팡매칭      제품 검색어별 쿠팡 판매 여부·가격·예상 수수료
    6_초안          업로드용 한국어 초안 전문

openpyxl 이 없으면 친절한 안내와 함께 RuntimeError.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, List, Sequence

FONT_NAME = "Arial"

# 훅 유형 판별 (draft.py 와 같은 기준, 사람이 읽는 라벨)
_HOOK_LABELS = (
    ("논쟁형", r"(unpopular opinion|hot take|controversial|change my mind)"),
    ("리스트형", r"\b\d+\s+(ways|things|tips|reasons|rules|steps|mistakes|habits|lessons)\b"),
    ("방법형", r"(how to|here'?s how|step by step)"),
    ("비밀/충격형", r"\b(secret|truth|nobody|never|stop doing|warning)\b"),
    ("질문형", r"\?\s*$"),
    ("스토리형", r"^\s*(i|my|we)\b"),
)

# 따라할 포인트 자동 메모
_TAKEAWAY_RULES = (
    (lambda p: (p.get("reply_rate_pct") or 0) >= 1.0, "댓글율 1%↑ — 질문/의견 유도 구조를 그대로 차용"),
    (lambda p: (p.get("share_rate_pct") or 0) >= 0.5, "공유율 높음 — 저장·공유할 정보(목록/수치)를 담은 글"),
    (lambda p: (p.get("reach_multiple") or 0) >= 5, "팔로워 대비 도달 5배↑ — 팔로워 없어도 먹히는 소재"),
    (lambda p: (p.get("views_per_hour") or 0) >= 5000, "확산 속도 빠름 — 지금 뜨는 주제"),
    (lambda p: len(p.get("text") or "") <= 120, "초단문 — 첫 줄에서 승부"),
    (lambda p: bool(p.get("has_media")), "사진/영상 동반"),
)


def hook_label(text: str) -> str:
    first = ""
    for line in (text or "").splitlines():
        if line.strip():
            first = line.strip()
            break
    for label, pattern in _HOOK_LABELS:
        if re.search(pattern, first, re.I | re.M):
            return label
    if re.search(r"\d", first):
        return "숫자형"
    return "일반"


def takeaways(post: Dict[str, Any]) -> str:
    hits = [note for rule, note in _TAKEAWAY_RULES if rule(post)]
    return " / ".join(hits[:3]) if hits else "상위권 진입 — 소재 자체를 참고"


def _require_openpyxl():
    try:
        import openpyxl  # noqa: WPS433
    except ImportError as exc:  # noqa: BLE001
        raise RuntimeError(
            "엑셀 출력에는 openpyxl 이 필요합니다:  pip install openpyxl"
        ) from exc
    return openpyxl


def write_xlsx(result: Dict[str, Any], path: str | Path) -> Path:
    """분석 결과 → 벤치마킹용 엑셀."""
    _require_openpyxl()
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    body_font = Font(name=FONT_NAME, size=10)
    title_font = Font(name=FONT_NAME, size=13, bold=True, color="1F3864")
    note_font = Font(name=FONT_NAME, size=9, italic=True, color="595959")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    band = PatternFill("solid", fgColor="F2F5FA")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")

    wb = Workbook()
    wb.remove(wb.active)

    def add_sheet(name: str, title: str, headers: Sequence[str], rows: Sequence[Sequence[Any]],
                  widths: Sequence[int], subtitle: str = "", wrap_cols: Sequence[int] = ()):
        ws = wb.create_sheet(name)
        ws["A1"] = title
        ws["A1"].font = title_font
        if subtitle:
            ws["A2"] = subtitle
            ws["A2"].font = note_font
        start = 3
        for c, head in enumerate(headers, 1):
            cell = ws.cell(row=start, column=c, value=head)
            cell.font, cell.fill, cell.border = head_font, head_fill, border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for r, row in enumerate(rows, start + 1):
            for c, value in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=value)
                cell.font, cell.border = body_font, border
                cell.alignment = wrap if c in wrap_cols else top
                if (r - start) % 2 == 0:
                    cell.fill = band
        for c, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.row_dimensions[start].height = 26
        ws.freeze_panes = ws.cell(row=start + 1, column=1)
        return ws, start

    posts = result.get("posts") or []
    summary = result.get("summary") or {}
    options = result.get("options") or {}
    analysis = result.get("analysis") or {}
    keywords = ", ".join(options.get("keywords") or []) or (options.get("from_json") or "-")

    # 1. 벤치마킹 -------------------------------------------------------
    rows = []
    for i, post in enumerate(posts, 1):
        coupang = post.get("coupang") or {}
        product = (coupang.get("product") or {}) if coupang else {}
        rows.append([
            i, post.get("score"), "@" + str(post.get("username") or ""),
            post.get("followers"), int(post.get("effective_views") or 0),
            post.get("likes"), post.get("replies"),
            (post.get("reposts") or 0) + (post.get("quotes") or 0) + (post.get("shares") or 0),
            (post.get("reply_rate_pct") or 0) / 100, (post.get("share_rate_pct") or 0) / 100,
            hook_label(post.get("text") or ""),
            (post.get("text") or "").replace("\r", " ").strip()[:600],
            takeaways(post),
            "판매중" if coupang.get("found") else ("없음" if coupang.get("keyword") else "-"),
            product.get("name") or "",
            post.get("url"),
        ])
    ws1, s1 = add_sheet(
        "1_벤치마킹", "해외 스레드 벤치마킹 보드",
        ["#", "점수", "계정", "팔로워", "노출", "좋아요", "댓글", "공유", "댓글율", "공유율",
         "훅 유형", "원문", "따라할 포인트", "쿠팡", "매칭 상품", "링크"],
        rows, [5, 7, 16, 11, 11, 10, 9, 9, 9, 9, 12, 60, 40, 8, 26, 44],
        subtitle=(f"키워드: {keywords} · 분석 {summary.get('count', 0)}건 중 상위 {len(posts)}건 · "
                  f"중앙 노출 {summary.get('median_views', 0):,}"),
        wrap_cols=(12, 13),
    )
    last1 = s1 + len(rows)
    for r in range(s1 + 1, last1 + 1):
        for c in (4, 5, 6, 7, 8):
            ws1.cell(row=r, column=c).number_format = "#,##0"
        for c in (9, 10):
            ws1.cell(row=r, column=c).number_format = "0.00%"
        link = ws1.cell(row=r, column=16)
        if link.value:
            link.hyperlink = link.value
            link.font = Font(name=FONT_NAME, size=10, color="0563C1", underline="single")
    if rows:
        avg_row = last1 + 1
        ws1.cell(row=avg_row, column=1, value="평균").font = Font(name=FONT_NAME, size=10, bold=True)
        for col, letter in ((2, "B"), (5, "E"), (6, "F"), (7, "G"), (8, "H"), (9, "I"), (10, "J")):
            cell = ws1.cell(row=avg_row, column=col,
                            value=f"=AVERAGE({letter}{s1+1}:{letter}{last1})")
            cell.font = Font(name=FONT_NAME, size=10, bold=True)
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
            cell.border = border
            cell.number_format = "0.00%" if col in (9, 10) else "#,##0"

    # 2. 패턴 -----------------------------------------------------------
    pattern_rows: List[List[Any]] = []
    for group_name, key in (("훅 유형", "hooks"), ("글 길이", "lengths"),
                            ("미디어", "media"), ("발행 시간대(KST)", "hours")):
        for g in analysis.get(key) or []:
            pattern_rows.append([
                group_name, g["name"], g["count"], g["avg_score"], g["lift"],
                g["median_views"], (g["avg_reply_rate_pct"] or 0) / 100,
                (g["avg_share_rate_pct"] or 0) / 100,
            ])
    ws2, s2 = add_sheet(
        "2_훅패턴", "무엇을 따라할까 — 패턴별 성과",
        ["구분", "유형", "글 수", "평균 점수", "배수(lift)", "중앙 노출", "평균 댓글율", "평균 공유율"],
        pattern_rows, [16, 40, 8, 11, 11, 12, 12, 12],
        subtitle="배수(lift) 1.0 초과 = 전체 평균보다 잘 된 유형. 이 유형을 우선 따라하세요.",
    )
    for r in range(s2 + 1, s2 + len(pattern_rows) + 1):
        ws2.cell(row=r, column=5).number_format = "0.00\"×\""
        ws2.cell(row=r, column=6).number_format = "#,##0"
        ws2.cell(row=r, column=7).number_format = "0.00%"
        ws2.cell(row=r, column=8).number_format = "0.00%"

    # 3. 키워드 ---------------------------------------------------------
    term_rows = [["단어", t["term"], t["top_count"], t["rest_count"], t["lift"], t["avg_score"]]
                 for t in analysis.get("keywords") or []]
    term_rows += [["해시태그", "#" + t["term"], t["top_count"], t["rest_count"], t["lift"], t["avg_score"]]
                  for t in analysis.get("hashtags") or []]
    ws3, s3 = add_sheet(
        "3_키워드", "상위권에서 유독 많이 나온 표현",
        ["구분", "표현", "상위권 등장", "하위권 등장", "lift", "평균 점수"],
        term_rows, [10, 26, 12, 12, 10, 11],
        subtitle="lift = 상위권 등장률 ÷ 하위권 등장률. 값이 클수록 '터진 글에만 있는' 표현.",
    )
    for r in range(s3 + 1, s3 + len(term_rows) + 1):
        ws3.cell(row=r, column=5).number_format = "0.00\"×\""

    # 4. 계정 -----------------------------------------------------------
    author_rows = [[
        "@" + a["username"], a["posts"], a["avg_score"], a["median_views"],
        a.get("followers"), a["best_url"],
    ] for a in analysis.get("authors") or []]
    ws4, s4 = add_sheet(
        "4_계정", "벤치마킹할 계정",
        ["계정", "수집된 글", "평균 점수", "중앙 노출", "팔로워", "대표 글"],
        author_rows, [22, 11, 11, 12, 12, 46],
        subtitle="평균 점수가 높은 계정 = 이 주제에서 꾸준히 먹히는 계정. 최근 글 전체를 훑어볼 대상.",
    )
    for r in range(s4 + 1, s4 + len(author_rows) + 1):
        for c in (4, 5):
            ws4.cell(row=r, column=c).number_format = "#,##0"
        link = ws4.cell(row=r, column=6)
        if link.value:
            link.hyperlink = link.value
            link.font = Font(name=FONT_NAME, size=10, color="0563C1", underline="single")

    # 5. 쿠팡 매칭 -------------------------------------------------------
    match_rows = []
    for m in result.get("matches") or []:
        product = m.get("product") or {}
        match_rows.append([
            m.get("keyword"), "판매중" if m.get("found") else "없음",
            product.get("name") or m.get("note") or "",
            product.get("price"), (product.get("commission_rate_pct") or 0) / 100,
            product.get("commission_per_sale"),
            "O" if product.get("is_rocket") else "",
            "O" if product.get("affiliate") else "",
            product.get("url") or "",
        ])
    ws5, s5 = add_sheet(
        "5_쿠팡매칭", "쿠팡 판매 매칭",
        ["검색어", "상태", "상품", "가격", "수수료율(추정)", "건당 수익(추정)", "로켓", "제휴링크", "링크"],
        match_rows, [18, 9, 44, 12, 14, 15, 7, 9, 44],
        subtitle="수수료율은 카테고리 평균 추정치입니다 — 실제 요율은 쿠팡 파트너스 공지를 확인하세요.",
    )
    for r in range(s5 + 1, s5 + len(match_rows) + 1):
        ws5.cell(row=r, column=4).number_format = '#,##0"원"'
        ws5.cell(row=r, column=5).number_format = "0.0%"
        ws5.cell(row=r, column=6).number_format = '#,##0"원"'
        link = ws5.cell(row=r, column=9)
        if link.value:
            link.hyperlink = link.value
            link.font = Font(name=FONT_NAME, size=10, color="0563C1", underline="single")

    # 6. 초안 -----------------------------------------------------------
    draft_rows = [[
        i, d.get("hook"), d.get("text"), d.get("char_count"),
        "예" if d.get("translated") else "아니오",
        " / ".join(d.get("warnings") or []), d.get("source_url"),
    ] for i, d in enumerate(result.get("drafts") or [], 1)]
    ws6, s6 = add_sheet(
        "6_초안", "업로드용 한국어 초안",
        ["#", "훅", "전문", "글자 수", "번역됨", "확인 사항", "원문"],
        draft_rows, [5, 34, 80, 9, 9, 40, 44],
        subtitle="그대로 올리지 말고 본인 말투로 고쳐 쓰세요. 쿠팡 파트너스 대가성 문구는 자동 포함되어 있습니다.",
        wrap_cols=(2, 3, 6),
    )
    for r in range(s6 + 1, s6 + len(draft_rows) + 1):
        ws6.row_dimensions[r].height = 150
        link = ws6.cell(row=r, column=7)
        if link.value:
            link.hyperlink = link.value
            link.font = Font(name=FONT_NAME, size=10, color="0563C1", underline="single")

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
